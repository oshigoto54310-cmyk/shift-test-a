"""出力テスト。出力にも権限スコープが効いていることを確認する。"""
from __future__ import annotations

import io

from openpyxl import load_workbook

from .conftest import create_order, future_date, get_products

DELIVERY = future_date(40)


def _setup_orders(store_a, store_b):
    a = get_products(store_a, "テストベンダーA")[0]
    b = get_products(store_a, "テストベンダーB")[0]
    create_order(store_a, [a["id"], b["id"]], confirm=True, delivery_date=DELIVERY)
    create_order(store_b, [a["id"], b["id"]], confirm=True, delivery_date=DELIVERY)


def _xlsx_text(content: bytes) -> str:
    wb = load_workbook(io.BytesIO(content))
    parts = []
    for ws in wb.worksheets:
        parts.append(ws.title)
        for row in ws.iter_rows(values_only=True):
            parts.append(" ".join(str(c) for c in row if c is not None))
    return "\n".join(parts)


def test_ベンダーは自社データのみ出力される(store_a, store_b, vendor_a):
    _setup_orders(store_a, store_b)

    res = vendor_a.get("/api/exports/orders", params={"fmt": "xlsx", "delivery_date": DELIVERY})
    assert res.status_code == 200
    text = _xlsx_text(res.content)
    assert "テストベンダーA" in text
    assert "テストベンダーB" not in text
    assert "テスト商品B1" not in text

    # CSV でも同じ
    res = vendor_a.get("/api/exports/orders", params={"fmt": "csv", "delivery_date": DELIVERY})
    assert res.status_code == 200
    body = res.content.decode("utf-8")
    assert "テストベンダーA" in body
    assert "テストベンダーB" not in body

    # 印刷用PDF（HTML）でも同じ
    res = vendor_a.get("/api/exports/orders", params={"fmt": "pdf", "delivery_date": DELIVERY})
    assert res.status_code == 200
    assert "テストベンダーB" not in res.text


def test_店舗は自店舗データのみ出力される(store_a, store_b):
    _setup_orders(store_a, store_b)

    res = store_a.get("/api/exports/orders", params={"fmt": "csv", "delivery_date": DELIVERY})
    assert res.status_code == 200
    body = res.content.decode("utf-8")
    assert "テスト店舗A" in body
    assert "テスト店舗B" not in body

    res = store_b.get("/api/exports/orders", params={"fmt": "csv", "delivery_date": DELIVERY})
    body = res.content.decode("utf-8")
    assert "テスト店舗B" in body
    assert "テスト店舗A" not in body


def test_本部は全データを出力できる(store_a, store_b, hq):
    _setup_orders(store_a, store_b)
    res = hq.get("/api/exports/orders", params={"fmt": "csv", "delivery_date": DELIVERY})
    assert res.status_code == 200
    body = res.content.decode("utf-8")
    for name in ("テスト店舗A", "テスト店舗B", "テストベンダーA", "テストベンダーB"):
        assert name in body


def test_本部はベンダーを切り替えて出力できる(store_a, store_b, hq):
    _setup_orders(store_a, store_b)
    vendors = hq.get("/api/vendors").json()
    vendor_b_id = next(v["id"] for v in vendors if v["name"] == "テストベンダーB")

    res = hq.get("/api/exports/orders", params={
        "fmt": "csv", "delivery_date": DELIVERY, "vendor_id": vendor_b_id,
    })
    body = res.content.decode("utf-8")
    assert "テストベンダーB" in body
    assert "テストベンダーA" not in body


def test_ベンダー別Excelはシートが分かれる(store_a, store_b, hq):
    _setup_orders(store_a, store_b)
    res = hq.get("/api/exports/orders", params={
        "fmt": "xlsx", "group_by": "vendor", "delivery_date": DELIVERY,
    })
    wb = load_workbook(io.BytesIO(res.content))
    assert set(wb.sheetnames) >= {"テストベンダーA", "テストベンダーB"}

    res = hq.get("/api/exports/orders", params={
        "fmt": "xlsx", "group_by": "store", "delivery_date": DELIVERY,
    })
    wb = load_workbook(io.BytesIO(res.content))
    assert set(wb.sheetnames) >= {"テスト店舗A", "テスト店舗B"}


def test_操作ログの出力は本部と管理者のみ(store_a, vendor_a, hq, admin):
    assert store_a.get("/api/exports/audit-logs").status_code == 403
    assert vendor_a.get("/api/exports/audit-logs").status_code == 403
    assert hq.get("/api/exports/audit-logs").status_code == 200
    assert admin.get("/api/exports/audit-logs").status_code == 200


def test_全ての出力形式が利用できる(hq):
    endpoints = [
        "orders", "product-summary", "delivery-date-summary", "shortages",
        "partials", "unanswered", "change-requests", "order-histories", "audit-logs",
    ]
    for ep in endpoints:
        for fmt in ("xlsx", "csv", "pdf"):
            res = hq.get(f"/api/exports/{ep}", params={"fmt": fmt})
            assert res.status_code == 200, f"{ep} / {fmt} が失敗: {res.text[:200]}"
            assert len(res.content) > 0


def test_出力操作が操作ログに残る(hq):
    hq.get("/api/exports/orders", params={"fmt": "xlsx"})
    hq.get("/api/exports/orders", params={"fmt": "csv"})
    hq.get("/api/exports/orders", params={"fmt": "pdf"})

    logs = hq.get("/api/histories/audit", params={"limit": 200}).json()
    actions = {log["action"] for log in logs}
    assert {"EXPORT_EXCEL", "EXPORT_CSV", "EXPORT_PDF"} <= actions


def test_商品別集約がベンダースコープで絞られる(store_a, store_b, vendor_b):
    _setup_orders(store_a, store_b)
    res = vendor_b.get("/api/exports/product-summary", params={"fmt": "csv", "delivery_date": DELIVERY})
    body = res.content.decode("utf-8")
    assert "テスト商品B1" in body
    assert "テスト商品A1" not in body
