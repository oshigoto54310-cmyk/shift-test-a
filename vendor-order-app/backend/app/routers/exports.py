"""出力機能（Excel / CSV / 印刷用PDF）。

出力にも必ず権限スコープを適用する。
ベンダーユーザーは自ベンダー分、店舗ユーザーは自店舗分のみ出力できる。
"""
from __future__ import annotations

import csv
import io
from datetime import date, datetime, timezone

from fastapi import APIRouter, Depends, HTTPException, Query, Request, Response, status
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, joinedload

from ..audit import log_action
from ..config import APP_TZ
from ..constants import (
    CHANGE_REQUEST_STATUS_LABELS,
    ORDER_STATUS_LABELS,
    RESPONSE_TYPE_LABELS,
    AuditAction,
    ChangeRequestStatus,
    OrderStatus,
    ResponseType,
)
from ..database import get_db
from ..deps import AccessScope, scope_dependency
from ..models import (
    AuditLog,
    ChangeRequest,
    Order,
    OrderHistory,
    OrderItem,
    Store,
    User,
    Vendor,
    VendorResponse,
)

router = APIRouter(prefix="/api/exports", tags=["出力"])

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _jst(dt: datetime | None) -> str:
    if dt is None:
        return ""
    return dt.replace(tzinfo=timezone.utc).astimezone(APP_TZ).strftime("%Y/%m/%d %H:%M")


def _status_label(code: str | None) -> str:
    if not code:
        return ""
    try:
        return ORDER_STATUS_LABELS.get(OrderStatus(code), code)
    except ValueError:
        return code


# --------------------------------------------------------------------------
# 共通: スコープ付きの明細クエリ
# --------------------------------------------------------------------------
def _item_query(
    db: Session,
    scope: AccessScope,
    *,
    delivery_date: date | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    statuses: list[str] | None = None,
):
    q = (
        db.query(OrderItem, Order, Store, Vendor)
        .join(Order, Order.id == OrderItem.order_id)
        .join(Store, Store.id == Order.store_id)
        .join(Vendor, Vendor.id == Order.vendor_id)
        .filter(Order.deleted_at.is_(None), OrderItem.deleted_at.is_(None))
    )
    # ---- 権限スコープ（ここを外すと他社データが漏れる）----
    if scope.store_id is not None:
        q = q.filter(Order.store_id == scope.store_id)
    if scope.vendor_id is not None:
        q = q.filter(Order.vendor_id == scope.vendor_id)

    if delivery_date:
        q = q.filter(Order.delivery_date == delivery_date)
    if date_from:
        q = q.filter(Order.delivery_date >= date_from)
    if date_to:
        q = q.filter(Order.delivery_date <= date_to)
    if statuses:
        q = q.filter(OrderItem.status.in_(statuses))
    return q.order_by(Order.delivery_date, Vendor.code, Store.code, OrderItem.id)


# --------------------------------------------------------------------------
# 共通: 出力ヘルパ
# --------------------------------------------------------------------------
def _autosize(ws) -> None:
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        length = max((len(str(c.value)) if c.value is not None else 0) for c in column_cells)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(length + 4, 10), 44)


def _build_sheet(ws, headers: list[str], rows: list[list]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    _autosize(ws)


def _xlsx_response(wb: Workbook, filename: str) -> Response:
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return Response(
        content=buffer.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _csv_response(headers: list[str], rows: list[list], filename: str) -> Response:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(headers)
    writer.writerows(rows)
    # Excel で開いたときに文字化けしないよう BOM 付き UTF-8 にする
    content = "﻿" + buffer.getvalue()
    return Response(
        content=content.encode("utf-8"),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _html_print_response(title: str, headers: list[str], rows: list[list], subtitle: str) -> Response:
    """印刷用HTML（ブラウザの印刷ダイアログからPDF保存する）。

    サーバー側で外部バイナリに依存せず、どの環境でもPDF化できる方式を採っている。
    """
    def esc(v) -> str:
        return (
            str(v if v is not None else "")
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

    thead = "".join(f"<th>{esc(h)}</th>" for h in headers)
    tbody = "".join(
        "<tr>" + "".join(f"<td>{esc(c)}</td>" for c in row) + "</tr>" for row in rows
    )
    html = f"""<!DOCTYPE html>
<html lang="ja"><head><meta charset="UTF-8">
<title>{esc(title)}</title>
<style>
 @page {{ size: A4 landscape; margin: 12mm; }}
 body {{ font-family: "Hiragino Sans", "Noto Sans JP", "Yu Gothic", sans-serif; color:#111; }}
 h1 {{ font-size: 16pt; margin: 0 0 4px; }}
 .sub {{ font-size: 9pt; color:#555; margin-bottom: 10px; }}
 table {{ border-collapse: collapse; width: 100%; font-size: 9pt; }}
 th, td {{ border: 1px solid #999; padding: 4px 6px; }}
 th {{ background: #1F3864; color: #fff; }}
 tr:nth-child(even) td {{ background: #f4f6fa; }}
 @media print {{ .noprint {{ display: none; }} }}
</style></head>
<body>
<h1>{esc(title)}</h1>
<div class="sub">{esc(subtitle)}</div>
<button class="noprint" onclick="window.print()">印刷 / PDF保存</button>
<table><thead><tr>{thead}</tr></thead><tbody>{tbody}</tbody></table>
</body></html>"""
    return Response(content=html, media_type="text/html; charset=utf-8")


def _deliver(
    fmt: str,
    *,
    title: str,
    sheet_name: str,
    headers: list[str],
    rows: list[list],
    filename_base: str,
    subtitle: str,
    db: Session,
    user: User,
    request: Request,
    export_kind: str,
):
    fmt = (fmt or "xlsx").lower()
    action = {
        "xlsx": AuditAction.EXPORT_EXCEL,
        "csv": AuditAction.EXPORT_CSV,
        "pdf": AuditAction.EXPORT_PDF,
    }.get(fmt)
    if action is None:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "形式は xlsx / csv / pdf のいずれかです")

    log_action(
        db, user, action, target_type="export", target_id=export_kind,
        detail={"種別": export_kind, "形式": fmt, "件数": len(rows)}, request=request, commit=True,
    )

    stamp = datetime.now(APP_TZ).strftime("%Y%m%d_%H%M")
    if fmt == "csv":
        return _csv_response(headers, rows, f"{filename_base}_{stamp}.csv")
    if fmt == "pdf":
        return _html_print_response(title, headers, rows, subtitle)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    _build_sheet(ws, headers, rows)
    return _xlsx_response(wb, f"{filename_base}_{stamp}.xlsx")


# --------------------------------------------------------------------------
# 発注（ベンダー別 / 店舗別）
# --------------------------------------------------------------------------
ORDER_HEADERS = [
    "納品日", "発注番号", "店舗コード", "店舗名", "ベンダーコード", "ベンダー名",
    "JANコード", "商品名", "規格", "発注単位", "ケース入数",
    "ケース数", "バラ数", "発注数量", "確定数量", "原価", "金額",
    "ステータス", "締め時間", "回答期限", "備考",
]


def _order_rows(query) -> list[list]:
    rows = []
    for item, order, store, vendor in query.all():
        cost = float(item.cost) if item.cost is not None else None
        rows.append([
            order.delivery_date.strftime("%Y/%m/%d"),
            order.order_no,
            store.code,
            store.name,
            vendor.code,
            vendor.name,
            item.jan_code,
            item.product_name,
            item.spec or "",
            item.order_unit,
            item.case_qty,
            item.qty_case,
            item.qty_loose,
            item.quantity,
            item.confirmed_quantity if item.confirmed_quantity is not None else "",
            cost if cost is not None else "",
            round(cost * item.quantity, 2) if cost is not None else "",
            _status_label(item.status),
            _jst(order.deadline_at),
            _jst(order.reply_deadline_at),
            item.note or "",
        ])
    return rows


@router.get("/orders")
def export_orders(
    request: Request,
    fmt: str = Query(default="xlsx", pattern="^(xlsx|csv|pdf)$"),
    group_by: str = Query(default="vendor", pattern="^(vendor|store|none)$"),
    delivery_date: date | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    db: Session = Depends(get_db),
    scope: AccessScope = Depends(scope_dependency),
):
    """発注一覧の出力。group_by=vendor でベンダー別、store で店舗別にシートを分ける。"""
    query = _item_query(
        db, scope, delivery_date=delivery_date, date_from=date_from, date_to=date_to
    )
    records = query.all()

    subtitle = f"出力日時 {datetime.now(APP_TZ):%Y/%m/%d %H:%M} / 件数 {len(records)}"
    stamp = datetime.now(APP_TZ).strftime("%Y%m%d_%H%M")

    if fmt in {"csv", "pdf"} or group_by == "none":
        rows = _order_rows(query)
        title = {"vendor": "ベンダー別発注一覧", "store": "店舗別発注一覧", "none": "発注一覧"}[group_by]
        return _deliver(
            fmt, title=title, sheet_name="発注一覧", headers=ORDER_HEADERS, rows=rows,
            filename_base="orders", subtitle=subtitle, db=db, user=scope.user,
            request=request, export_kind=f"発注一覧({group_by})",
        )

    # Excel: グループごとにシートを分ける
    grouped: dict[str, list[list]] = {}
    for item, order, store, vendor in records:
        key = vendor.name if group_by == "vendor" else store.name
        cost = float(item.cost) if item.cost is not None else None
        grouped.setdefault(key, []).append([
            order.delivery_date.strftime("%Y/%m/%d"), order.order_no, store.code, store.name,
            vendor.code, vendor.name, item.jan_code, item.product_name, item.spec or "",
            item.order_unit, item.case_qty, item.qty_case, item.qty_loose, item.quantity,
            item.confirmed_quantity if item.confirmed_quantity is not None else "",
            cost if cost is not None else "",
            round(cost * item.quantity, 2) if cost is not None else "",
            _status_label(item.status), _jst(order.deadline_at), _jst(order.reply_deadline_at),
            item.note or "",
        ])

    wb = Workbook()
    wb.remove(wb.active)
    if not grouped:
        _build_sheet(wb.create_sheet("データなし"), ORDER_HEADERS, [])
    for key, rows in sorted(grouped.items()):
        safe = key.replace("/", "／").replace("\\", "＼")[:31] or "無題"
        _build_sheet(wb.create_sheet(safe), ORDER_HEADERS, rows)

    log_action(db, scope.user, AuditAction.EXPORT_EXCEL, target_type="export",
               target_id=f"発注一覧({group_by})",
               detail={"形式": "xlsx", "件数": len(records)}, request=request, commit=True)
    prefix = "orders_by_vendor" if group_by == "vendor" else "orders_by_store"
    return _xlsx_response(wb, f"{prefix}_{stamp}.xlsx")


# --------------------------------------------------------------------------
# 商品別集約 / 納品日別集約
# --------------------------------------------------------------------------
@router.get("/product-summary")
def export_product_summary(
    request: Request,
    fmt: str = Query(default="xlsx", pattern="^(xlsx|csv|pdf)$"),
    delivery_date: date | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    db: Session = Depends(get_db),
    scope: AccessScope = Depends(scope_dependency),
):
    """商品別集約。ベンダーが仕入・出荷を組むための集計。"""
    records = _item_query(
        db, scope, delivery_date=delivery_date, date_from=date_from, date_to=date_to
    ).all()

    agg: dict[tuple, dict] = {}
    for item, order, store, vendor in records:
        key = (vendor.id, item.product_id, order.delivery_date)
        row = agg.setdefault(key, {
            "delivery_date": order.delivery_date, "vendor_name": vendor.name,
            "jan_code": item.jan_code, "product_name": item.product_name,
            "spec": item.spec or "", "order_unit": item.order_unit, "case_qty": item.case_qty,
            "qty_case": 0, "qty_loose": 0, "quantity": 0, "store_count": set(), "confirmed": 0,
        })
        row["qty_case"] += item.qty_case
        row["qty_loose"] += item.qty_loose
        row["quantity"] += item.quantity
        row["confirmed"] += item.confirmed_quantity or 0
        row["store_count"].add(store.id)

    headers = ["納品日", "ベンダー名", "JANコード", "商品名", "規格", "発注単位",
               "ケース入数", "合計ケース数", "合計バラ数", "合計発注数量", "合計確定数量", "発注店舗数"]
    rows = [
        [r["delivery_date"].strftime("%Y/%m/%d"), r["vendor_name"], r["jan_code"],
         r["product_name"], r["spec"], r["order_unit"], r["case_qty"],
         r["qty_case"], r["qty_loose"], r["quantity"], r["confirmed"], len(r["store_count"])]
        for r in sorted(agg.values(), key=lambda x: (x["delivery_date"], x["vendor_name"], x["product_name"]))
    ]
    return _deliver(
        fmt, title="商品別集約", sheet_name="商品別集約", headers=headers, rows=rows,
        filename_base="product_summary",
        subtitle=f"出力日時 {datetime.now(APP_TZ):%Y/%m/%d %H:%M} / 商品数 {len(rows)}",
        db=db, user=scope.user, request=request, export_kind="商品別集約",
    )


@router.get("/delivery-date-summary")
def export_delivery_date_summary(
    request: Request,
    fmt: str = Query(default="xlsx", pattern="^(xlsx|csv|pdf)$"),
    date_from: date | None = None,
    date_to: date | None = None,
    db: Session = Depends(get_db),
    scope: AccessScope = Depends(scope_dependency),
):
    records = _item_query(db, scope, date_from=date_from, date_to=date_to).all()

    agg: dict[tuple, dict] = {}
    for item, order, store, vendor in records:
        key = (order.delivery_date, vendor.id, store.id)
        row = agg.setdefault(key, {
            "delivery_date": order.delivery_date, "vendor_name": vendor.name,
            "store_name": store.name, "orders": set(), "items": 0, "quantity": 0,
            "shortage": 0, "partial": 0, "pending": 0,
        })
        row["orders"].add(order.id)
        row["items"] += 1
        row["quantity"] += item.quantity
        if item.status == str(OrderStatus.SHORTAGE):
            row["shortage"] += 1
        elif item.status == str(OrderStatus.PARTIAL):
            row["partial"] += 1
        elif item.status == str(OrderStatus.VENDOR_PENDING):
            row["pending"] += 1

    headers = ["納品日", "ベンダー名", "店舗名", "発注件数", "明細数", "合計数量",
               "欠品明細数", "一部納品明細数", "未回答明細数"]
    rows = [
        [r["delivery_date"].strftime("%Y/%m/%d"), r["vendor_name"], r["store_name"],
         len(r["orders"]), r["items"], r["quantity"], r["shortage"], r["partial"], r["pending"]]
        for r in sorted(agg.values(), key=lambda x: (x["delivery_date"], x["vendor_name"], x["store_name"]))
    ]
    return _deliver(
        fmt, title="納品日別集約", sheet_name="納品日別集約", headers=headers, rows=rows,
        filename_base="delivery_date_summary",
        subtitle=f"出力日時 {datetime.now(APP_TZ):%Y/%m/%d %H:%M}",
        db=db, user=scope.user, request=request, export_kind="納品日別集約",
    )


# --------------------------------------------------------------------------
# 欠品 / 一部納品 / 未回答
# --------------------------------------------------------------------------
def _response_detail_rows(db: Session, records) -> list[list]:
    item_ids = [item.id for item, *_ in records]
    responses = {}
    if item_ids:
        for r in (
            db.query(VendorResponse)
            .filter(VendorResponse.order_item_id.in_(item_ids), VendorResponse.is_latest.is_(True))
            .all()
        ):
            responses[r.order_item_id] = r

    rows = []
    for item, order, store, vendor in records:
        r = responses.get(item.id)
        rows.append([
            order.delivery_date.strftime("%Y/%m/%d"), order.order_no, store.name, vendor.name,
            item.jan_code, item.product_name, item.spec or "", item.quantity,
            (r.deliverable_qty if r and r.deliverable_qty is not None else ""),
            (r.shortage_qty if r and r.shortage_qty is not None else ""),
            (RESPONSE_TYPE_LABELS.get(ResponseType(r.response_type), r.response_type) if r else ""),
            (r.reason or r.shortage_reason or "") if r else "",
            (r.next_available_date.strftime("%Y/%m/%d") if r and r.next_available_date else ""),
            (r.sub_product_name or "") if r else "",
            _jst(r.responded_at) if r else "",
            _status_label(item.status),
        ])
    return rows


SHORTAGE_HEADERS = ["納品日", "発注番号", "店舗名", "ベンダー名", "JANコード", "商品名", "規格",
                    "発注数量", "納品可能数量", "不足数量", "回答区分", "理由",
                    "次回納品可能日", "代替商品名", "回答日時", "ステータス"]


@router.get("/shortages")
def export_shortages(
    request: Request,
    fmt: str = Query(default="xlsx", pattern="^(xlsx|csv|pdf)$"),
    delivery_date: date | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    db: Session = Depends(get_db),
    scope: AccessScope = Depends(scope_dependency),
):
    records = _item_query(
        db, scope, delivery_date=delivery_date, date_from=date_from, date_to=date_to,
        statuses=[str(OrderStatus.SHORTAGE)],
    ).all()
    return _deliver(
        fmt, title="欠品一覧", sheet_name="欠品一覧", headers=SHORTAGE_HEADERS,
        rows=_response_detail_rows(db, records), filename_base="shortages",
        subtitle=f"出力日時 {datetime.now(APP_TZ):%Y/%m/%d %H:%M} / 件数 {len(records)}",
        db=db, user=scope.user, request=request, export_kind="欠品一覧",
    )


@router.get("/partials")
def export_partials(
    request: Request,
    fmt: str = Query(default="xlsx", pattern="^(xlsx|csv|pdf)$"),
    delivery_date: date | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    db: Session = Depends(get_db),
    scope: AccessScope = Depends(scope_dependency),
):
    records = _item_query(
        db, scope, delivery_date=delivery_date, date_from=date_from, date_to=date_to,
        statuses=[str(OrderStatus.PARTIAL)],
    ).all()
    return _deliver(
        fmt, title="一部納品一覧", sheet_name="一部納品一覧", headers=SHORTAGE_HEADERS,
        rows=_response_detail_rows(db, records), filename_base="partial_deliveries",
        subtitle=f"出力日時 {datetime.now(APP_TZ):%Y/%m/%d %H:%M} / 件数 {len(records)}",
        db=db, user=scope.user, request=request, export_kind="一部納品一覧",
    )


@router.get("/unanswered")
def export_unanswered(
    request: Request,
    fmt: str = Query(default="xlsx", pattern="^(xlsx|csv|pdf)$"),
    delivery_date: date | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    db: Session = Depends(get_db),
    scope: AccessScope = Depends(scope_dependency),
):
    """未回答一覧（ベンダー未確認のまま残っている明細）。"""
    records = _item_query(
        db, scope, delivery_date=delivery_date, date_from=date_from, date_to=date_to,
        statuses=[str(OrderStatus.VENDOR_PENDING)],
    ).all()
    headers = ["納品日", "発注番号", "店舗名", "ベンダー名", "JANコード", "商品名", "規格",
               "発注数量", "締め時間", "回答期限", "期限超過"]
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    rows = [
        [order.delivery_date.strftime("%Y/%m/%d"), order.order_no, store.name, vendor.name,
         item.jan_code, item.product_name, item.spec or "", item.quantity,
         _jst(order.deadline_at), _jst(order.reply_deadline_at),
         "超過" if order.reply_deadline_at and order.reply_deadline_at < now else ""]
        for item, order, store, vendor in records
    ]
    return _deliver(
        fmt, title="未回答一覧", sheet_name="未回答一覧", headers=headers, rows=rows,
        filename_base="unanswered",
        subtitle=f"出力日時 {datetime.now(APP_TZ):%Y/%m/%d %H:%M} / 件数 {len(rows)}",
        db=db, user=scope.user, request=request, export_kind="未回答一覧",
    )


# --------------------------------------------------------------------------
# 変更申請 / 変更履歴
# --------------------------------------------------------------------------
@router.get("/change-requests")
def export_change_requests(
    request: Request,
    fmt: str = Query(default="xlsx", pattern="^(xlsx|csv|pdf)$"),
    db: Session = Depends(get_db),
    scope: AccessScope = Depends(scope_dependency),
):
    q = (
        db.query(ChangeRequest, Order, Store, Vendor)
        .join(Order, Order.id == ChangeRequest.order_id)
        .join(Store, Store.id == Order.store_id)
        .join(Vendor, Vendor.id == Order.vendor_id)
        .filter(Order.deleted_at.is_(None))
    )
    if scope.store_id is not None:
        q = q.filter(Order.store_id == scope.store_id)
    if scope.vendor_id is not None:
        q = q.filter(Order.vendor_id == scope.vendor_id)
    records = q.order_by(ChangeRequest.requested_at.desc()).all()

    item_ids = {cr.order_item_id for cr, *_ in records}
    items = {
        i.id: i for i in db.query(OrderItem).filter(OrderItem.id.in_(item_ids)).all()
    } if item_ids else {}
    user_ids = {cr.requester_id for cr, *_ in records} | {cr.approver_id for cr, *_ in records}
    users = {
        u.id: u.name for u in db.query(User).filter(User.id.in_({u for u in user_ids if u})).all()
    } if user_ids else {}

    headers = ["申請日時", "納品日", "発注番号", "店舗名", "ベンダー名", "商品名",
               "変更前ケース", "変更前バラ", "変更前数量", "申請ケース", "申請バラ", "申請数量",
               "変更理由", "状態", "申請者", "承認者", "承認日時", "却下理由", "ベンダー確認日時"]
    rows = [
        [_jst(cr.requested_at), order.delivery_date.strftime("%Y/%m/%d"), order.order_no,
         store.name, vendor.name, items.get(cr.order_item_id).product_name if items.get(cr.order_item_id) else "",
         cr.before_case, cr.before_loose, cr.before_quantity,
         cr.requested_case, cr.requested_loose, cr.requested_quantity,
         cr.reason,
         CHANGE_REQUEST_STATUS_LABELS.get(ChangeRequestStatus(cr.status), cr.status),
         users.get(cr.requester_id, ""), users.get(cr.approver_id, ""),
         _jst(cr.approved_at), cr.reject_reason or "", _jst(cr.vendor_confirmed_at)]
        for cr, order, store, vendor in records
    ]
    return _deliver(
        fmt, title="変更申請一覧", sheet_name="変更申請一覧", headers=headers, rows=rows,
        filename_base="change_requests",
        subtitle=f"出力日時 {datetime.now(APP_TZ):%Y/%m/%d %H:%M} / 件数 {len(rows)}",
        db=db, user=scope.user, request=request, export_kind="変更申請一覧",
    )


@router.get("/order-histories")
def export_order_histories(
    request: Request,
    fmt: str = Query(default="xlsx", pattern="^(xlsx|csv|pdf)$"),
    db: Session = Depends(get_db),
    scope: AccessScope = Depends(scope_dependency),
):
    """発注変更履歴の出力。"""
    order_q = db.query(Order.id).filter(Order.deleted_at.is_(None))
    if scope.store_id is not None:
        order_q = order_q.filter(Order.store_id == scope.store_id)
    if scope.vendor_id is not None:
        order_q = order_q.filter(Order.vendor_id == scope.vendor_id)
    allowed = [row[0] for row in order_q.all()]

    histories = (
        db.query(OrderHistory)
        .filter(OrderHistory.order_id.in_(allowed))
        .order_by(OrderHistory.changed_at.desc())
        .limit(5000)
        .all()
        if allowed
        else []
    )
    orders = {
        o.id: o
        for o in db.query(Order)
        .options(joinedload(Order.store), joinedload(Order.vendor))
        .filter(Order.id.in_({h.order_id for h in histories}))
        .all()
    } if histories else {}
    items = {
        i.id: i
        for i in db.query(OrderItem)
        .filter(OrderItem.id.in_({h.order_item_id for h in histories if h.order_item_id}))
        .all()
    } if histories else {}
    users = {
        u.id: u.name
        for u in db.query(User)
        .filter(User.id.in_({h.changed_by for h in histories if h.changed_by}))
        .all()
    } if histories else {}

    headers = ["変更日時", "発注番号", "店舗名", "ベンダー名", "商品名",
               "変更前ケース", "変更後ケース", "変更前バラ", "変更後バラ",
               "変更前数量", "変更後数量", "変更理由", "変更者", "締め前/締め後",
               "申請者ID", "承認者ID", "承認日時", "却下理由", "ベンダー確認日時"]
    rows = []
    for h in histories:
        order = orders.get(h.order_id)
        item = items.get(h.order_item_id)
        rows.append([
            _jst(h.changed_at), order.order_no if order else "",
            order.store.name if order and order.store else "",
            order.vendor.name if order and order.vendor else "",
            item.product_name if item else "",
            h.case_before, h.case_after, h.loose_before, h.loose_after,
            h.qty_before, h.qty_after, h.reason or "", users.get(h.changed_by, ""),
            "締め後" if h.is_after_deadline else "締め前",
            h.requester_id or "", h.approver_id or "", _jst(h.approved_at),
            h.reject_reason or "", _jst(h.vendor_confirmed_at),
        ])
    return _deliver(
        fmt, title="発注変更履歴", sheet_name="発注変更履歴", headers=headers, rows=rows,
        filename_base="order_histories",
        subtitle=f"出力日時 {datetime.now(APP_TZ):%Y/%m/%d %H:%M} / 件数 {len(rows)}",
        db=db, user=scope.user, request=request, export_kind="発注変更履歴",
    )


@router.get("/audit-logs")
def export_audit_logs(
    request: Request,
    fmt: str = Query(default="xlsx", pattern="^(xlsx|csv|pdf)$"),
    date_from: date | None = None,
    date_to: date | None = None,
    db: Session = Depends(get_db),
    scope: AccessScope = Depends(scope_dependency),
):
    """操作ログの出力。全社の記録なので本部・管理者のみ。"""
    if not scope.is_cross_tenant:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "操作ログを出力する権限がありません")

    q = db.query(AuditLog)
    if date_from:
        q = q.filter(AuditLog.created_at >= datetime.combine(date_from, datetime.min.time()))
    if date_to:
        q = q.filter(AuditLog.created_at <= datetime.combine(date_to, datetime.max.time()))
    rows_db = q.order_by(AuditLog.created_at.desc()).limit(10000).all()

    headers = ["日時", "ユーザーID", "メールアドレス", "権限", "操作", "対象種別", "対象ID", "詳細", "IPアドレス"]
    rows = [
        [_jst(r.created_at), r.user_id or "", r.user_email or "", r.role_code or "",
         r.action, r.target_type or "", r.target_id or "", r.detail or "", r.ip_address or ""]
        for r in rows_db
    ]
    return _deliver(
        fmt, title="操作ログ", sheet_name="操作ログ", headers=headers, rows=rows,
        filename_base="audit_logs",
        subtitle=f"出力日時 {datetime.now(APP_TZ):%Y/%m/%d %H:%M} / 件数 {len(rows)}",
        db=db, user=scope.user, request=request, export_kind="操作ログ",
    )
